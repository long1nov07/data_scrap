"""So sánh các mã CÙNG NGÀNH.

Đọc dữ liệu từ các file Excel đã tạo trong output/HOSE và output/HNX
(giá trị lưu dạng số nên đọc lại chính xác, không cần gọi lại API), nhóm theo
ngành (sector) và xuất 1 file Excel:
  - Sheet "Mục lục": danh sách ngành + số mã + vài chỉ số trung vị.
  - Mỗi ngành 1 sheet: mỗi dòng 1 mã, xếp theo vốn hoá, kèm dòng "Trung vị ngành".

Dùng:
    python compare.py                       # đọc ./output, xuất output/So_sanh_nganh.xlsx
    python compare.py --output D:/du_lieu    # thư mục khác
"""
from __future__ import annotations

import argparse
import glob
import os
import re
import statistics
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Tái dùng style & định dạng số cho đồng bộ với file gốc
from excel_builder import (
    _BORDER, _CENTER, _FMT_NUM, _FMT_PCT, _FMT_VND, _LEFT, _RIGHT,
    _LABEL_FONT, _NAVY, _TITLE_FONT, _fit_columns, _style_header_row,
)

# Tên ngành tiếng Việt (sector do vnstock trả về là tiếng Anh)
SECTOR_VI = {
    "Construction & Materials": "Xây dựng & Vật liệu",
    "Industrial Goods & Services": "Hàng & DV Công nghiệp",
    "Real Estate": "Bất động sản",
    "Utilities": "Tiện ích (Điện/Nước/Khí)",
    "Food & Beverage": "Thực phẩm & Đồ uống",
    "Basic Resources": "Tài nguyên cơ bản",
    "Chemicals": "Hoá chất",
    "Financial Services": "Dịch vụ tài chính",
    "Personal & Household Goods": "Hàng tiêu dùng & Gia dụng",
    "Health Care": "Chăm sóc sức khoẻ",
    "Banks": "Ngân hàng",
    "Media": "Truyền thông",
    "Technology": "Công nghệ",
    "Travel & Leisure": "Du lịch & Giải trí",
    "Retail": "Bán lẻ",
    "Automobiles & Parts": "Ô tô & Phụ tùng",
    "Insurance": "Bảo hiểm",
    "Oil & Gas": "Dầu khí",
    "Telecommunications": "Viễn thông",
}

# Cột so sánh: (nhãn, khoá đọc, sheet nguồn, định dạng, chia_tỷ)
# sheet nguồn: 'ov' = Tổng quan, 'kpi' = bảng chỉ tiêu then chốt (trong Tổng quan),
#              'rt' = Chỉ số tài chính. Giá trị lấy ở cột kỳ gần nhất (cột B).
COLUMNS = [
    ("Mã", None, None, None, False),
    ("Sàn", None, None, None, False),
    # Định giá
    ("Vốn hoá (tỷ)", "Vốn hoá (VND)", "ov", _FMT_VND, True),
    ("Giá (VND)", "Giá hiện tại (VND)", "ov", _FMT_VND, False),
    ("P/E", "P/E (TTM)", "ov", _FMT_NUM, False),
    ("P/B", "P/B", "ov", _FMT_NUM, False),
    # Sinh lời
    ("ROE quý", "ROE (quý)", "rt", _FMT_PCT, False),
    ("ROA quý", "ROA (quý)", "rt", _FMT_PCT, False),
    ("Biên LN gộp", "Biên LN gộp", "rt", _FMT_PCT, False),
    ("Biên LN ròng", "Biên LN ròng", "rt", _FMT_PCT, False),
    # Quy mô & tăng trưởng
    ("Doanh thu quý (tỷ)", "Doanh thu thuần", "kpi", _FMT_VND, True),
    ("LNST quý (tỷ)", "LN thuần sau thuế", "kpi", _FMT_VND, True),
    ("Tăng trưởng DT QoQ", "Tăng trưởng doanh thu (QoQ)", "rt", _FMT_PCT, False),
    # An toàn tài chính
    ("Nợ/VCSH (D/E)", "Nợ phải trả / VCSH (D/E)", "rt", _FMT_NUM, False),
    ("TT hiện hành", "Hệ số thanh toán hiện hành", "rt", _FMT_NUM, False),
]
# Cột số (bỏ Mã, Sàn) để tính trung vị & canh format
_NUM_COLS = COLUMNS[2:]
SORT_KEY = "Vốn hoá (VND)"  # xếp mã trong ngành theo vốn hoá giảm dần


def _rows_map(ws):
    """Gom các dòng của sheet thành dict {nhãn cột A: tuple giá trị cả dòng}."""
    m = {}
    for row in ws.iter_rows(values_only=True):
        if row and row[0] is not None:
            key = str(row[0]).strip()
            if key and key not in m:  # giữ dòng khớp ĐẦU TIÊN
                m[key] = row
    return m


def read_company(path, exchange):
    """Đọc 1 file mã -> dict {nhãn cột: giá trị}. None nếu lỗi."""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return None
    try:
        ov = _rows_map(wb["Tổng quan"])
        rt = _rows_map(wb["Chỉ số tài chính"])
    finally:
        wb.close()

    src = {"ov": ov, "kpi": ov, "rt": rt}  # kpi cũng nằm trong sheet Tổng quan
    sym = os.path.splitext(os.path.basename(path))[0]

    _sec_row = ov.get("Ngành (sector)")
    rec = {"Mã": sym, "Sàn": exchange,
           "_sector": _sec_row[1] if _sec_row and len(_sec_row) > 1 else None}
    for label, key, sheet, _fmt, scale in COLUMNS:
        if key is None:
            continue
        row = src[sheet].get(key)
        val = row[1] if row and len(row) > 1 else None
        if isinstance(val, (int, float)) and scale:
            val = val / 1e9
        rec[label] = val if isinstance(val, (int, float)) else None
    rec[SORT_KEY] = (ov.get("Vốn hoá (VND)") or (None, None))[1] \
        if ov.get("Vốn hoá (VND)") else None
    return rec


def collect(out_dir):
    """Quét output/HOSE và output/HNX (KHÔNG đệ quy -> bỏ qua thư mục người dùng tự tạo)."""
    records = []
    for exch in ("HOSE", "HNX"):
        for path in glob.glob(os.path.join(out_dir, exch, "*.xlsx")):
            if os.path.basename(path).startswith("~$"):
                continue
            rec = read_company(path, exch)
            if rec:
                records.append(rec)
    return records


def _sanitize_sheet_name(name, used):
    bad = r'[\\/*?:\[\]]'
    name = re.sub(bad, "-", name)[:31].strip()
    base, i = name, 1
    while name.lower() in used:
        suffix = f" ({i})"
        name = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(name.lower())
    return name


def _median(vals):
    nums = [v for v in vals if isinstance(v, (int, float))]
    return statistics.median(nums) if nums else None


def _write_industry_sheet(wb, sheet_name, sector_en, rows):
    ws = wb.create_sheet(sheet_name)
    headers = [c[0] for c in COLUMNS]

    ws.cell(1, 1, f"{SECTOR_VI.get(sector_en, sector_en)} — {len(rows)} mã").font = _TITLE_FONT
    ws.cell(2, 1, "Xếp theo vốn hoá giảm dần • đơn vị tiền: tỷ đồng (trừ Giá là VND)").font = \
        Font(italic=True, size=9, color="808080")

    hr = 4
    for c, h in enumerate(headers, 1):
        ws.cell(hr, c, h)
    _style_header_row(ws, hr, len(headers))

    rows = sorted(rows, key=lambda r: (r.get(SORT_KEY) is None, -(r.get(SORT_KEY) or 0)))
    for i, rec in enumerate(rows):
        r = hr + 1 + i
        for c, (label, _k, _s, fmt, _sc) in enumerate(COLUMNS, 1):
            v = rec.get(label)
            cell = ws.cell(r, c, v)
            if label == "Mã":
                cell.font = Font(bold=True)
                cell.alignment = _LEFT
            elif label == "Sàn":
                cell.alignment = _CENTER
            else:
                cell.alignment = _RIGHT
                if isinstance(v, (int, float)):
                    cell.number_format = fmt
            cell.border = _BORDER

    # Dòng trung vị ngành
    mr = hr + 1 + len(rows)
    mc = ws.cell(mr, 1, "TRUNG VỊ NGÀNH")
    mc.font = _LABEL_FONT
    for c, (label, _k, _s, fmt, _sc) in enumerate(COLUMNS, 1):
        if label in ("Mã", "Sàn"):
            continue
        med = _median([rec.get(label) for rec in rows])
        cell = ws.cell(mr, c, med)
        if isinstance(med, (int, float)):
            cell.number_format = fmt
        cell.font = Font(bold=True)
        cell.alignment = _RIGHT
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
        cell.border = _BORDER
    ws.cell(mr, 1).fill = PatternFill("solid", fgColor="D9E1F2")

    ws.freeze_panes = f"B{hr + 1}"
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(headers))}{mr - 1}"
    _fit_columns(ws, first_col_min=8, first_col_max=12, other_min=11, other_max=20)
    return ws


def _write_index_sheet(wb, groups):
    ws = wb.create_sheet("Mục lục")
    ws.cell(1, 1, "SO SÁNH CÁC MÃ CÙNG NGÀNH").font = _TITLE_FONT
    ws.cell(2, 1, "Nhấp tên ngành ở các tab bên dưới. Trung vị = giá trị đại diện ngành.").font = \
        Font(italic=True, size=9, color="808080")

    headers = ["Ngành", "Số mã", "Trung vị P/E", "Trung vị P/B", "Trung vị ROE quý", "Trung vị D/E"]
    hr = 4
    for c, h in enumerate(headers, 1):
        ws.cell(hr, c, h)
    _style_header_row(ws, hr, len(headers))

    order = sorted(groups.items(), key=lambda kv: -len(kv[1]))
    for i, (sector_en, rows) in enumerate(order):
        r = hr + 1 + i
        ws.cell(r, 1, SECTOR_VI.get(sector_en, sector_en)).alignment = _LEFT
        ws.cell(r, 2, len(rows)).alignment = _CENTER
        for c, label, fmt in [(3, "P/E", _FMT_NUM), (4, "P/B", _FMT_NUM),
                              (5, "ROE quý", _FMT_PCT), (6, "Nợ/VCSH (D/E)", _FMT_NUM)]:
            med = _median([rec.get(label) for rec in rows])
            cell = ws.cell(r, c, med)
            if isinstance(med, (int, float)):
                cell.number_format = fmt
            cell.alignment = _RIGHT
        for c in range(1, len(headers) + 1):
            ws.cell(r, c).border = _BORDER
    ws.freeze_panes = "A5"
    _fit_columns(ws, first_col_min=22, first_col_max=30, other_min=13, other_max=18)
    return ws


def build(out_dir, out_file):
    records = collect(out_dir)
    if not records:
        print("Không tìm thấy file mã nào trong", out_dir)
        return
    # Nhóm theo sector
    groups = {}
    for rec in records:
        sector = rec.get("_sector") or "Chưa phân loại"
        groups.setdefault(sector, []).append(rec)

    wb = Workbook()
    wb.remove(wb.active)
    _write_index_sheet(wb, groups)

    used = {"mục lục"}
    for sector_en, rows in sorted(groups.items(), key=lambda kv: -len(kv[1])):
        name = _sanitize_sheet_name(SECTOR_VI.get(sector_en, sector_en), used)
        _write_industry_sheet(wb, name, sector_en, rows)

    wb.save(out_file)
    print(f"Đã tạo: {out_file}")
    print(f"  Tổng số mã: {len(records)} | Số ngành: {len(groups)}")
    for sector_en, rows in sorted(groups.items(), key=lambda kv: -len(kv[1])):
        print(f"    {len(rows):4}  {SECTOR_VI.get(sector_en, sector_en)}")


def main():
    p = argparse.ArgumentParser(description="So sánh các mã cùng ngành từ file đã tạo.")
    p.add_argument("--output", default="output", help="Thư mục chứa HOSE/ và HNX/")
    p.add_argument("--out-file", default=None, help="Đường dẫn file Excel kết quả")
    args = p.parse_args()
    out_dir = os.path.abspath(args.output)
    out_file = args.out_file or os.path.join(out_dir, "So_sanh_nganh.xlsx")
    build(out_dir, out_file)


if __name__ == "__main__":
    main()
