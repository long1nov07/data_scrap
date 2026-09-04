"""Dựng file Excel (5 sheet) cho 1 doanh nghiệp từ dữ liệu đã xử lý."""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import analysis
import config

# --- Bảng màu / style ---
_NAVY = "1F3864"
_BLUE = "2F5496"
_LIGHT = "D9E1F2"
_GREY = "F2F2F2"
_WHITE = "FFFFFF"

_HDR_FONT = Font(bold=True, color=_WHITE, size=11)
_TITLE_FONT = Font(bold=True, color=_NAVY, size=14)
_LABEL_FONT = Font(bold=True, color=_NAVY, size=10)
_BORDER = Border(*[Side(style="thin", color="BFBFBF")] * 4)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_RIGHT = Alignment(horizontal="right", vertical="center")

_FMT_VND = "#,##0;[Red]-#,##0"
_FMT_PCT = "0.0%;[Red]-0.0%"
_FMT_NUM = "#,##0.00"


def _style_header_row(ws, row, ncols, fill=_BLUE):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = _HDR_FONT
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = _CENTER if c > 1 else _LEFT
        cell.border = _BORDER


def _display_len(cell) -> int:
    """Ước lượng số ký tự Excel cần để hiển thị 1 ô (kể cả dấu phân tách nghìn).

    Con số VND có thể tới hàng nghìn tỷ -> phải tính theo chuỗi ĐÃ định dạng,
    nếu không cột sẽ hẹp và Excel hiện '####'.
    """
    v = cell.value
    if v is None:
        return 0
    if isinstance(v, bool):
        return len(str(v))
    if isinstance(v, (int, float)):
        nf = cell.number_format or ""
        if "%" in nf:
            return len(f"{v * 100:,.1f}") + 1          # vd '-12.3%'
        if '"' in nf or "0.00" in nf:
            return len(f"{v:,.2f}") + 5                # số thập phân + hậu tố ' lần'
        return len(f"{v:,.0f}") + 1                     # số nguyên có dấu phẩy + dấu âm
    return len(str(v))


def _fit_columns(ws, first_col_min=24, first_col_max=52, other_min=12, other_max=30):
    """Canh độ rộng cột theo nội dung thực tế (tránh hiện '####')."""
    for col in range(1, ws.max_column + 1):
        longest = max(
            (_display_len(ws.cell(row=r, column=col)) for r in range(1, ws.max_row + 1)),
            default=0,
        )
        if col == 1:
            width = min(first_col_max, max(first_col_min, longest + 1))
        else:
            width = min(other_max, max(other_min, longest + 2))
        ws.column_dimensions[get_column_letter(col)].width = width


def _write_statement_sheet(wb, title, df, note="Đơn vị: VND"):
    """Sheet báo cáo: tiêu đề + ghi chú + bảng (Chỉ tiêu + kỳ) + cột %QoQ."""
    ws = wb.create_sheet(title[:31])
    periods = [c for c in df.columns if c != "Chỉ tiêu"]

    ws.cell(row=1, column=1, value=title).font = _TITLE_FONT
    ws.cell(row=2, column=1, value=note).font = Font(italic=True, size=9, color="808080")

    header_row = 4
    headers = ["Chỉ tiêu"] + periods + (["%QoQ"] if len(periods) >= 2 else [])
    for c, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=c, value=h)
    _style_header_row(ws, header_row, len(headers))

    for r, (_, row) in enumerate(df.iterrows(), start=header_row + 1):
        name = str(row["Chỉ tiêu"])
        is_total = name.isupper()  # dòng tổng (VD: TỔNG CỘNG TÀI SẢN)
        c1 = ws.cell(row=r, column=1, value=name)
        c1.alignment = _LEFT
        c1.border = _BORDER
        if is_total:
            c1.font = _LABEL_FONT
        vals = [row[p] for p in periods]
        for c, v in enumerate(vals, start=2):
            cell = ws.cell(row=r, column=c, value=None if pd.isna(v) else float(v))
            cell.number_format = _FMT_VND
            cell.alignment = _RIGHT
            cell.border = _BORDER
            if is_total:
                cell.font = Font(bold=True)
        if len(periods) >= 2:
            qoq = analysis._growth(
                None if pd.isna(vals[0]) else vals[0],
                None if pd.isna(vals[1]) else vals[1],
            )
            cell = ws.cell(row=r, column=len(headers),
                           value=None if qoq is None else qoq)
            cell.number_format = _FMT_PCT
            cell.alignment = _RIGHT
            cell.border = _BORDER
        if is_total:
            for c in range(1, len(headers) + 1):
                ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=_LIGHT)

    ws.freeze_panes = f"B{header_row + 1}"
    _fit_columns(ws)
    return ws


def _write_overview_sheet(wb, fin, val):
    ws = wb.create_sheet(config.SHEET_OVERVIEW[:31])
    ov = fin["overview"]
    sym = fin["symbol"]

    ws.cell(row=1, column=1, value=f"{sym} — {ov.get('organ_name') or ''}").font = _TITLE_FONT
    ws.merge_cells("A1:D1")

    info = [
        ("Sàn niêm yết", ov.get("com_group_code")),
        ("Ngành (sector)", ov.get("sector")),
        ("Loại hình", "Ngân hàng" if ov.get("is_bank") else "Phi tài chính"),
        ("Giá hiện tại (VND)", ov.get("current_price")),
        ("Vốn hoá (VND)", ov.get("market_cap")),
        ("Số CP lưu hành", ov.get("issue_share")),
        ("P/E (TTM)", val.get("pe_ttm")),
        ("P/B", val.get("pb")),
        ("LNST cổ đông mẹ TTM (VND)", val.get("ln_ttm")),
        ("Khuyến nghị", ov.get("rating")),
        ("Giá mục tiêu (VND)", ov.get("target_price")),
        ("Sở hữu nước ngoài", ov.get("foreigner_percentage")),
        ("Ngày niêm yết", str(ov.get("listing_date") or "")[:10]),
    ]
    r = 3
    for label, value in info:
        lc = ws.cell(row=r, column=1, value=label)
        lc.font = _LABEL_FONT
        lc.fill = PatternFill("solid", fgColor=_GREY)
        lc.border = _BORDER
        vc = ws.cell(row=r, column=2, value=value)
        vc.border = _BORDER
        if isinstance(value, (int, float)):
            if "%" in label or "nước ngoài" in label:
                vc.number_format = _FMT_PCT
            elif label in ("P/E (TTM)", "P/B"):
                vc.number_format = _FMT_NUM
            else:
                vc.number_format = _FMT_VND
        r += 1

    # Bảng chỉ tiêu then chốt
    r += 1
    ws.cell(row=r, column=1, value="CHỈ TIÊU THEN CHỐT QUA CÁC KỲ").font = _LABEL_FONT
    r += 1
    ki = analysis.key_indicators(fin)
    header_row = r
    for c, h in enumerate(ki.columns, start=1):
        ws.cell(row=header_row, column=c, value=h)
    _style_header_row(ws, header_row, len(ki.columns), fill=_NAVY)
    for _, row in ki.iterrows():
        r += 1
        kc = ws.cell(row=r, column=1, value=str(row["Chỉ tiêu"]))
        kc.border = _BORDER
        kc.alignment = _LEFT
        for c, col in enumerate(ki.columns[1:], start=2):
            v = row[col]
            cell = ws.cell(row=r, column=c, value=None if pd.isna(v) else float(v))
            cell.number_format = _FMT_PCT if col == "%QoQ" else _FMT_VND
            cell.alignment = _RIGHT
            cell.border = _BORDER

    ws.freeze_panes = "A2"
    _fit_columns(ws, first_col_min=28, first_col_max=40, other_min=16, other_max=32)
    return ws


def _write_ratios_sheet(wb, fin):
    ws = wb.create_sheet(config.SHEET_RATIOS[:31])
    df = analysis.financial_ratios(fin)
    fmt = df.attrs.get("fmt", [])
    periods = [c for c in df.columns if c != "Chỉ số"]

    ws.cell(row=1, column=1, value=config.SHEET_RATIOS).font = _TITLE_FONT
    ws.cell(row=2, column=1,
            value="Chỉ số tự tính từ BCTC (theo quý, chưa quy đổi năm)").font = \
        Font(italic=True, size=9, color="808080")

    header_row = 4
    for c, h in enumerate(df.columns, start=1):
        ws.cell(row=header_row, column=c, value=h)
    _style_header_row(ws, header_row, len(df.columns))

    for idx, (_, row) in enumerate(df.iterrows()):
        r = header_row + 1 + idx
        pct, suffix = fmt[idx] if idx < len(fmt) else (False, "")
        lc = ws.cell(row=r, column=1, value=str(row["Chỉ số"]))
        lc.border = _BORDER
        lc.alignment = _LEFT
        for c, p in enumerate(periods, start=2):
            v = row[p]
            cell = ws.cell(row=r, column=c, value=None if pd.isna(v) else float(v))
            if pct:
                cell.number_format = _FMT_PCT
            elif suffix:
                cell.number_format = f'0.00"{suffix}"'
            else:
                cell.number_format = _FMT_NUM
            cell.alignment = _RIGHT
            cell.border = _BORDER

    ws.freeze_panes = f"B{header_row + 1}"
    _fit_columns(ws, first_col_min=30, first_col_max=42, other_min=12, other_max=20)
    return ws


def build_company_workbook(fin: dict, out_path: str):
    """Dựng & lưu file Excel 5 sheet cho 1 mã."""
    val = analysis.valuation_snapshot(fin)
    wb = Workbook()
    wb.remove(wb.active)  # bỏ sheet mặc định

    _write_overview_sheet(wb, fin, val)
    _write_statement_sheet(wb, config.SHEET_BALANCE, fin["balance"])
    _write_statement_sheet(wb, config.SHEET_INCOME, fin["income"])
    _write_statement_sheet(wb, config.SHEET_CASHFLOW, fin["cashflow"])
    _write_ratios_sheet(wb, fin)

    wb.save(out_path)
    return out_path
